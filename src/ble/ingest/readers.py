"""Read tender BoQs and bids into lines with evidence, choosing the cheapest route that works:
GAEB XML is parsed, spreadsheets are read cell by cell, PDFs with a text layer are read from it,
and only pages without one are sent to OCR."""

from __future__ import annotations

import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
from dataclasses import asdict
from pathlib import Path

from openpyxl import load_workbook

from ble import gaeb
from ble.config import DATA, load
from ble.ingest.repair import repair
from ble.ingest.model import BidDoc, BidLine, Evidence, Sentence, TenderDoc, TenderPosition, de_num
from ble.ingest.table import Word, letter_sentences, lines_of, parse_table


def doc_id(path: Path, root: Path) -> str:
    h = hashlib.sha256(path.read_bytes()).hexdigest()[:12]
    return f"{path.relative_to(root).as_posix()}@{h}"


def read_tender(path: Path, root: Path, version: str) -> TenderDoc:
    d = gaeb.read(path.read_bytes())
    did = doc_id(path, root)
    pos = [TenderPosition(it.oz, it.oz.split(".")[0], it.short, it.long, it.qty, it.unit, Evidence(did, xml_line=it.line, text=it.short)) for it in d.items]
    tid = path.parent.parent.name
    return TenderDoc(tid, version, d.boq_date, d.boq_label, d.project, d.categories, pos, did)


# -- PDFs ----------------------------------------------------------------------------------------
def pdf_words(path: Path) -> tuple[list[Word], int, list[int]]:
    """Words from the text layer, the page count, and the pages that have no text layer."""
    import pdfplumber

    words: list[Word] = []
    empty: list[int] = []
    with pdfplumber.open(str(path)) as pdf:
        n = len(pdf.pages)
        for i, page in enumerate(pdf.pages, start=1):
            ws = page.extract_words(keep_blank_chars=False, use_text_flow=False)
            if len(ws) < 5:
                empty.append(i)
            words += [Word(w["text"], float(w["x0"]), float(w["x1"]), float(w["top"]), float(w["bottom"]), i) for w in ws]
    return words, n, empty


def _tessdata() -> str | None:
    for d in (os.environ.get("TESSDATA_PREFIX"), "/usr/share/tesseract-ocr/5/tessdata", "/usr/share/tesseract-ocr/4.00/tessdata", "/usr/share/tessdata"):
        if d and Path(d, "deu.traineddata").exists():
            return d
    return None


def ocr_words(path: Path, pages: list[int]) -> list[Word]:
    """OCR the given pages (1-based) with tesseract; boxes are converted to PDF points."""
    import pypdfium2 as pdfium

    cfg = load("ingest")
    dpi, psm, lang = cfg["routes"]["ocr"]["dpi"], cfg["routes"]["ocr"]["psm"], cfg["routes"]["ocr"]["languages"]
    min_conf = cfg["ocr"]["min_word_confidence"]
    if not shutil.which("tesseract"):
        raise RuntimeError("tesseract is not installed")
    td = _tessdata()
    doc = pdfium.PdfDocument(str(path))
    out: list[Word] = []
    k = 72.0 / dpi
    sha = hashlib.sha256(path.read_bytes()).hexdigest()[:16]
    cache_dir = Path(os.environ.get("BLE_CACHE", DATA / ".cache")) / "ocr"
    with tempfile.TemporaryDirectory() as tmp:
        for p in pages:
            key = cache_dir / f"{sha}-p{p}-{dpi}-{psm}-{lang if td else 'eng'}.json"
            if key.exists():  # same file, page and settings: OCR once (re-ingestion is idempotent and free)
                out += [Word(**w) for w in json.loads(key.read_text())]
                continue
            page_words: list[Word] = []
            img = doc[p - 1].render(scale=dpi / 72).to_pil().convert("L")
            f = Path(tmp) / f"p{p}.png"
            img.save(f)
            cmd = ["tesseract", str(f), "stdout", "-l", lang if td else "eng", "--psm", str(psm), "tsv"]
            if td:
                cmd[1:1] = ["--tessdata-dir", td]
            res = subprocess.run(cmd, capture_output=True, text=True, check=True)
            for row in res.stdout.splitlines()[1:]:
                c = row.split("\t")
                if len(c) < 12 or not c[11].strip():
                    continue
                conf = float(c[10])
                if conf < min_conf:
                    continue
                x, y, w, h = (int(c[6]), int(c[7]), int(c[8]), int(c[9]))
                page_words.append(Word(c[11], x * k, (x + w) * k, y * k, (y + h) * k, p, conf))
            key.parent.mkdir(parents=True, exist_ok=True)
            key.write_text(json.dumps([asdict(w) for w in page_words]))
            out += page_words
    return out


# -- spreadsheets --------------------------------------------------------------------------------
SYN = {
    "no": {"oz", "pos.", "pos", "position", "nr.", "nr"},
    "text": {"kurztext", "bezeichnung", "leistung", "text"},
    "qty": {"menge", "anzahl"},
    "unit": {"einheit", "me", "einh."},
    "ep": {"ep", "ep (€)", "einheitspreis", "ep netto"},
    "gp": {"gp", "gp (€)", "gesamtpreis", "gp netto"},
    "product": {"fabrikat", "angebotenes fabrikat"},
}


def xlsx_lines(path: Path, did: str) -> tuple[list[BidLine], float | None, list[str]]:
    wb = load_workbook(io.BytesIO(path.read_bytes()), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    cols: dict[str, int] = {}
    out: list[BidLine] = []
    trade, total, warnings = "", None, []
    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        vals = ["" if v is None else v for v in row]
        low = [str(v).strip().lower() for v in vals]
        if not cols:
            found = {k: next((i for i, t in enumerate(low) if t in names), -1) for k, names in SYN.items()}
            if found["text"] >= 0 and found["qty"] >= 0 and found["ep"] >= 0:
                cols = {k: v for k, v in found.items() if v >= 0}
            continue
        g = lambda k: vals[cols[k]] if k in cols and cols[k] < len(vals) else ""  # noqa: E731
        text = str(g("text")).strip()
        no = str(g("no")).strip()
        ev = Evidence(did, row=r, text=" | ".join(str(v) for v in vals if v != ""))
        if "angebotssumme" in text.lower():
            total = _f(g("gp"))
            continue
        if text.lower().startswith("summe"):
            continue
        unit, qty = str(g("unit")).strip(), _f(g("qty"))
        if not unit and qty is None and re.fullmatch(r"\d{1,2}", no):
            trade = no.zfill(2)
            continue
        if not text:
            continue
        ep_raw = g("ep")
        included = isinstance(ep_raw, str) and ep_raw.strip().lower().startswith("enth")
        ep, gp = (None if included else _f(ep_raw)), _f(g("gp"))
        product = str(g("product")).strip() or None if "product" in cols else None
        out.append(BidLine(len(out), no, trade, text, "", qty, unit, ep, gp, included or (ep == 0 and not gp), product, ev))
    if not cols:
        warnings.append("no header row found")
    return out, total, warnings


def _f(v: object) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return de_num(str(v))


# -- bids ----------------------------------------------------------------------------------------
def read_bid(bid_dir: Path, root: Path) -> BidDoc:
    t0 = time.perf_counter()
    cfg = load("ingest")["routes"]
    bid_id, tid = bid_dir.name, bid_dir.parent.parent.name
    files = {p.name: p for p in bid_dir.iterdir() if p.is_file()}
    lines: list[BidLine] = []
    letter: list[Sentence] = []
    total, label, company, warnings = None, "", "", []
    doc_ids: list[str] = []
    pages, cost, route = 0, 0.0, ""
    if "anschreiben.pdf" in files:
        did = doc_id(files["anschreiben.pdf"], root)
        doc_ids.append(did)
        words, n, empty = pdf_words(files["anschreiben.pdf"])
        if empty:
            words += ocr_words(files["anschreiben.pdf"], empty)
            cost += cfg["ocr"]["eur_per_page"] * len(empty)
        cost += cfg["pdf_text"]["eur_per_page"] * (n - len(empty))
        pages += n
        letter, total = letter_sentences(words, did)
        company = _company(words)
    main = next((files[n] for n in ("angebot.x84", "angebot.xlsx", "angebot.pdf", "angebot_scan.pdf") if n in files), None)
    if main is None:
        raise FileNotFoundError(f"no bid document in {bid_dir}")
    did = doc_id(main, root)
    doc_ids.append(did)
    if main.suffix == ".x84":
        route = "gaeb"
        d = gaeb.read(main.read_bytes())
        label, company = d.boq_label, d.company or company
        for it in d.items:
            ev = Evidence(did, xml_line=it.line, text=it.short)
            prod = next((v for k, v in it.bidder_texts.items() if "fabrikat" in k.lower()), None)
            lines.append(BidLine(len(lines), it.oz, it.oz.split(".")[0], it.short, it.long, it.qty, it.unit, it.up, it.it, it.up == 0 and it.it == 0, prod, ev))
        pages += 1
        cost += cfg["gaeb"]["eur_per_page"]
    elif main.suffix == ".xlsx":
        route = "xlsx"
        lines, t2, w = xlsx_lines(main, did)
        total = t2 if t2 is not None else total
        warnings += w
        pages += 1
        cost += cfg["xlsx"]["eur_per_page"]
    else:
        words, n, empty = pdf_words(main)
        route = "ocr" if len(empty) == n else "pdf_text"
        if empty:
            words = [w for w in words if w.page not in empty] + ocr_words(main, empty)
        cost += cfg["ocr"]["eur_per_page"] * len(empty) + cfg["pdf_text"]["eur_per_page"] * (n - len(empty))
        pages += n
        lines, t2, w = parse_table(words, did, scale_tol=1.4 if route == "ocr" else 1.0)
        warnings += w
        if route == "ocr":
            k = repair(lines)
            if k:
                warnings.append(f"{k} OCR'd lines corrected by their arithmetic")
        total = t2 if t2 is not None else total
        if not letter:
            letter, t3 = letter_sentences([x for x in words if x.page == 1], did)
            total = total if total is not None else t3
        company = company or _company([x for x in words if x.page == 1])
    return BidDoc(bid_id, tid, company, route, lines, letter, total, label, doc_ids, pages, round(cost, 6), round(time.perf_counter() - t0, 3), warnings)


def _company(words: list[Word]) -> str:
    """The letterhead: the first line of the first page."""
    lines = lines_of([w for w in words if w.page == 1])
    if not lines:
        return ""
    return " ".join(w.text for w in lines[0]).replace("Angebot ", "").strip()


def read_tender_versions(tender_dir: Path, root: Path) -> list[TenderDoc]:
    """All issued versions of a tender's BoQ, oldest first (by the BoQ date, then file name)."""
    files = sorted((tender_dir / "tender").glob("*.x83"))
    docs = [read_tender(f, root, "") for f in files]
    docs.sort(key=lambda d: (d.date, d.doc_id))
    for k, d in enumerate(docs, start=1):
        d.version = f"v{k}"
    return docs
